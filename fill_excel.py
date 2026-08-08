"""
기존 엑셀 파일(종목코드/기업명/연도/매출액/유형자산/연구개발비/직원수/산업 형식)에서
매출액/유형자산/연구개발비/직원수 칸을 DART Open API로 채워 넣는다.

기본은 비어 있는 칸만 채우고, 이미 값이 있는 칸은 건드리지 않는다.
--force 를 주면 이미 값이 있어도 API로 다시 받아서 덮어쓴다("N/A"라고 적힌,
그 해에 회사가 아직 없었다는 표시는 --force 에서도 건드리지 않는다).

연구개발비는 사업보고서 본문에서 "연구개발비용 계" 행을 직접 찾아서 읽는
방식이라(정형 API가 없음) 문서 포맷에 따라 못 찾을 수 있다 — 못 찾으면
그 칸은 그냥 비워 둔다. 산업/종목코드는 이 스크립트가 다루지 않는다.

사용법:
    set DART_API_KEY=발급받은키
    python fill_excel.py 입력파일.xlsx [출력파일.xlsx] [--force]

출력파일을 생략하면 "입력파일_filled.xlsx" 로 저장한다 (원본은 건드리지 않음).
"""

import sys

from openpyxl import load_workbook

from dart_financial_data import (
    fetch_employee_count,
    fetch_financials,
    fetch_rnd_expense,
    find_corp_code,
    load_corp_code_xml,
)

# 시트에 등장하는 기업명이 DART 공식 회사명과 다르면 여기에 매핑을 추가한다.
NAME_ALIASES = {
    "두산퓨어셀": "두산퓨얼셀",  # 정식 사명은 "두산퓨얼셀"
}

COL_CORP_CODE = 1
COL_NAME = 2
COL_YEAR = 3
COL_REVENUE = 4
COL_TANGIBLE = 5
COL_RND = 6
COL_EMPLOYEES = 7
COL_INDUSTRY = 8


def main():
    args = sys.argv[1:]
    force = "--force" in args
    args = [a for a in args if a != "--force"]

    if not args:
        sys.exit("사용법: python fill_excel.py 입력파일.xlsx [출력파일.xlsx] [--force]")
    in_path = args[0]
    out_path = args[1] if len(args) > 1 else in_path.rsplit(".", 1)[0] + "_filled.xlsx"

    wb = load_workbook(in_path)
    ws = wb.active

    # 기업명은 각 회사 블록의 첫 행에만 적혀 있고 그 아래는 비어 있으므로 채워 내려간다.
    current_name = None
    rows = []
    for r in range(2, ws.max_row + 1):
        name_cell = ws.cell(r, COL_NAME).value
        if name_cell:
            current_name = name_cell.strip()
        year = ws.cell(r, COL_YEAR).value
        if current_name and year:
            rows.append((r, current_name, int(year)))

    xml_text = load_corp_code_xml()
    corp_code_cache = {}

    def get_corp_code(name):
        if name not in corp_code_cache:
            lookup_name = NAME_ALIASES.get(name, name)
            corp_code_cache[name] = find_corp_code(xml_text, lookup_name)
        return corp_code_cache[name]

    financials_cache = {}

    # "N/A"는 그 해에 회사가 존재하지 않았다는 의도적인 표시이므로 --force 에서도 건드리지 않는다.
    def needs(value):
        if value == "N/A":
            return False
        if force:
            return True
        return value in (None, "")

    filled = 0
    for r, name, year in rows:
        need_revenue = needs(ws.cell(r, COL_REVENUE).value)
        need_tangible = needs(ws.cell(r, COL_TANGIBLE).value)
        need_rnd = needs(ws.cell(r, COL_RND).value)
        need_employees = needs(ws.cell(r, COL_EMPLOYEES).value)

        if not (need_revenue or need_tangible or need_rnd or need_employees):
            continue

        try:
            corp_code = get_corp_code(name)
        except ValueError as e:
            print(f"[건너뜀] {name} {year}: {e}")
            continue

        row_filled = 0

        if need_revenue or need_tangible:
            key = (corp_code, year)
            if key not in financials_cache:
                financials_cache[key] = fetch_financials(corp_code, year, name=name)
            revenue, tangible, fs_div = financials_cache[key]
            if need_revenue and revenue is not None:
                ws.cell(r, COL_REVENUE).value = revenue
                row_filled += 1
            if need_tangible and tangible is not None:
                ws.cell(r, COL_TANGIBLE).value = tangible
                row_filled += 1

        if need_rnd:
            rnd = fetch_rnd_expense(corp_code, year, name=name)
            if rnd is not None:
                ws.cell(r, COL_RND).value = rnd
                row_filled += 1

        if need_employees:
            employees = fetch_employee_count(corp_code, year, name=name)
            if employees is not None:
                ws.cell(r, COL_EMPLOYEES).value = employees
                row_filled += 1

        if row_filled:
            print(f"{name} {year} 처리 완료 ({row_filled}칸)")
        else:
            print(f"[값 없음] {name} {year}: DART에 해당 연도 데이터가 없어 그대로 둠")
        filled += row_filled

    wb.save(out_path)
    print(f"\n{filled}개 칸을 채웠습니다. 저장 위치: {out_path}")


if __name__ == "__main__":
    main()
